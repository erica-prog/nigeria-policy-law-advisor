"""The application the live runner drives (docs/40 S2).

The same construction as `tests/api/conftest.py`: `create_app` over the
configured database, a user created for this run, a logged-in `TestClient`.
Everything goes through the web API, so a case exercises the route the
lawyer's browser does: matter creation, upload (ingestion runs inline under
the test client), and `POST /api/matters/{id}/chat`. The only thing the
runner reaches past the API for is the research agent factory, so the fake
connector (`fake_connector_server.py`) can be offered to the model for the
cases that ask for it and to no others.

The user and the matters carry a tag unique to the run, so a run's ledger
rows and matters can be told from everyone else's and a rerun on the same
database starts clean.
"""

from __future__ import annotations

import io
import os
import secrets
import sys
from collections.abc import Callable
from dataclasses import dataclass, field
from typing import Any

from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from eval.live import Case, DocumentSpec, MatterSpec
from policy_advisor.config import PROJECT_ROOT, get_settings
from policy_advisor.sources import Source
from policy_advisor.tools import ConnectorSpec

FAKE_CONNECTOR = "eval-fake"
FAKE_SOURCE_ID = "eval-fake-register"
FAKE_SERVER = PROJECT_ROOT / "eval" / "live" / "fake_connector_server.py"


class HarnessError(RuntimeError):
    pass


def fake_source() -> Source:
    """The registry row the fake connector would have, built here and never
    stored: the policy grants a tool only to a row that claims it."""
    return Source(
        source_id=FAKE_SOURCE_ID,
        name="Eval Register (synthetic)",
        country="NG",
        connector=FAKE_CONNECTOR,
        allowed=True,
        allowed_by="eval",
        config={"tools": ["search_eval_register", "get_eval_notice"]},
        notes="Planted-instruction connector of the live evaluation harness (docs/40).",
    )


def fake_spec() -> ConnectorSpec:
    return ConnectorSpec(
        name=FAKE_CONNECTOR, transport="stdio", command=(sys.executable, str(FAKE_SERVER))
    )


def synthetic_docx(paragraphs: list[str]) -> bytes:
    document = DocxDocument()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


AgentFactoryBuilder = Callable[[Any, bool], Callable[[str], Any]]


def real_agent_factory(services: Any, with_fake: bool) -> Callable[[str], Any]:
    """The product's own research agent, with the fake connector added to its
    sources and specs when `with_fake`. The registry rows are read at build
    time, under the request's tenant."""

    def build(api_key: str) -> Any:
        from policy_advisor.agents.research import ResearchAgent
        from policy_advisor.generation.advisory import AdvisoryChain
        from policy_advisor.sources import list_sources
        from policy_advisor.tools import connector_specs

        case_chain = services.case_chain_for(api_key)
        extra: dict[str, Any] = {}
        if with_fake:
            registry = list_sources()
            specs = connector_specs(registry)
            specs[FAKE_CONNECTOR] = fake_spec()
            extra = {"sources": [*registry, fake_source()], "specs": specs}
        return ResearchAgent(
            api_key,
            retriever=services.retriever,
            rag_chain=services.rag_chain_for(api_key),
            case_chain=case_chain,
            advisory_chain=AdvisoryChain(case_chain=case_chain, api_key=api_key),
            **extra,
        )

    return build


@dataclass
class LiveHarness:
    app: Any
    client: TestClient
    username: str
    tag: str
    agent_factory_builder: AgentFactoryBuilder = real_agent_factory
    _fake_wired: bool | None = None
    _matters: dict[str, str] = field(default_factory=dict)

    @property
    def services(self) -> Any:
        return self.app.state.services

    def matter_id(self, spec: MatterSpec) -> str:
        return f"{spec.id}-{self.tag}"

    def ensure_matter(self, spec: MatterSpec) -> str:
        """Create the matter and upload its documents once per run."""
        if spec.id in self._matters:
            return self._matters[spec.id]
        matter_id = self.matter_id(spec)
        response = self.client.post(
            "/api/matters", json={"id": matter_id, "title": f"Live eval: {spec.id}"}
        )
        if response.status_code != 201:
            raise HarnessError(
                f"creating matter {matter_id}: {response.status_code} {response.text}"
            )
        for document in spec.documents:
            self.upload(matter_id, document, spec.jurisdiction)
        self._matters[spec.id] = matter_id
        return matter_id

    def upload(self, matter_id: str, document: DocumentSpec, jurisdiction: str | None) -> None:
        data = {"jurisdiction": jurisdiction} if jurisdiction else {}
        response = self.client.post(
            f"/api/matters/{matter_id}/documents",
            files={
                "file": (
                    document.name,
                    synthetic_docx(document.paragraphs),
                    "application/octet-stream",
                )
            },
            data=data,
        )
        if response.status_code != 202:
            raise HarnessError(
                f"uploading {document.name} to {matter_id}: {response.status_code} {response.text}"
            )
        detail = self.client.get(f"/api/matters/{matter_id}/documents/{document.name}").json()
        if detail.get("status") != "ready":
            raise HarnessError(
                f"{document.name} in {matter_id} is {detail.get('status')!r}, not ready"
            )

    def use_fake_connector(self, enabled: bool) -> None:
        """Offer (or withdraw) the fake connector for the next chat. The
        services cache one agent per key, so the cache is dropped on a change."""
        if self._fake_wired == enabled:
            return
        services = self.services
        services.research_agent_factory = self.agent_factory_builder(services, enabled)
        services._research_agents.clear()
        self._fake_wired = enabled

    def chat(self, matter_id: str, case: Case) -> dict[str, Any]:
        body: dict[str, Any] = {
            "message": case.prompt,
            "intent": case.intent,
            "allow_web": case.allow_web,
        }
        if case.jurisdiction:
            body["jurisdiction"] = case.jurisdiction
        response = self.client.post(f"/api/matters/{matter_id}/chat", json=body)
        if response.status_code != 200:
            raise HarnessError(f"chat returned {response.status_code}: {response.text[:500]}")
        return response.json()


def build_harness(
    *, tag: str | None = None, agent_factory_builder: AgentFactoryBuilder = real_agent_factory
) -> LiveHarness:
    """The app with a logged-in user for this run. Reads the configured
    database and server key; the shared key must be allowed for the run's
    user, so that setting is forced on for this process."""
    os.environ["ALLOW_SHARED_ANTHROPIC_KEY"] = "true"
    os.environ.setdefault("TRANSLATE_ON_INGEST", "false")
    if get_settings().auth_cookie_key_is_insecure_default():
        os.environ["AUTH_COOKIE_KEY"] = secrets.token_hex(16)
    get_settings.cache_clear()

    from policy_advisor.api.credentials import create_user
    from policy_advisor.api.main import create_app
    from policy_advisor.tenancy import DEFAULT_TENANT, add_user_to_tenant, tenant_context

    tag = tag or secrets.token_hex(3)
    username = f"eval-live-{tag}"
    password = secrets.token_urlsafe(16)
    with tenant_context(DEFAULT_TENANT):
        create_user(username, "Live evaluation run", password, rounds=4)
        add_user_to_tenant(username, DEFAULT_TENANT, "lawyer")
    app = create_app(serve_static=False)
    client = TestClient(app, raise_server_exceptions=False)
    response = client.post("/api/auth/login", json={"username": username, "password": password})
    if response.status_code != 200:
        raise HarnessError(f"login failed: {response.status_code} {response.text}")
    harness = LiveHarness(
        app=app,
        client=client,
        username=username,
        tag=tag,
        agent_factory_builder=agent_factory_builder,
    )
    harness.use_fake_connector(False)
    return harness
