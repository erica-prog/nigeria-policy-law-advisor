"""Rendering (docs/35): one outline, two files.

DOCX goes through docxtpl: a Word template with Jinja tags, the packaged
default or the firm's own, filled from the outline and the manifest. PPTX
is built with python-pptx from the same outline: a title slide, one slide
per section with the paragraphs as bullets and their sources in the notes,
and a closing slide of limitations and provenance counts. Rejected
paragraphs are replaced by a marker in both; unverified ones are marked.
Nothing is sent anywhere: both return bytes for the blob store.
"""

import io
from importlib import resources
from typing import Any

from docxtpl import DocxTemplate
from pptx import Presentation
from pptx.util import Inches, Pt

from policy_advisor.artefacts.outline import Outline
from policy_advisor.artefacts.provenance import Manifest

REJECTED_MARKER = (
    "[Paragraph removed: it cited case evidence as legal authority. See the provenance manifest.]"
)
UNVERIFIED_MARKER = " [UNVERIFIED: the cited source was not among the passages retrieved]"


def _cite(paragraph_sources: list[str]) -> str:
    return "; ".join(paragraph_sources)


def _paragraph_views(outline: Outline, manifest: Manifest) -> list[dict[str, Any]]:
    rejected = manifest.rejected_paragraphs
    unverified = manifest.unverified_paragraphs
    sections = []
    for s_index, section in enumerate(outline.sections):
        paragraphs = []
        for p_index, paragraph in enumerate(section.paragraphs):
            key = (s_index, p_index)
            if key in rejected:
                paragraphs.append({"text": REJECTED_MARKER, "cite": "", "status": "rejected"})
            else:
                text = paragraph.text + (UNVERIFIED_MARKER if key in unverified else "")
                paragraphs.append(
                    {
                        "text": text,
                        "cite": _cite(paragraph.sources),
                        "status": "unverified" if key in unverified else "ok",
                    }
                )
        sections.append({"heading": section.heading, "paragraphs": paragraphs})
    return sections


def _provenance_lines(manifest: Manifest) -> list[str]:
    lines = []
    for entry in manifest.entries:
        if entry.status == "instruction":
            continue
        lines.append(
            f"§{entry.section + 1}.{entry.paragraph + 1} {entry.pinpoint or '(no source)'}: {entry.status}"
            + (f"; {entry.instrument}" if entry.instrument else "")
            + (f"; version {entry.version_date}" if entry.version_date else "")
            + (f"; {entry.limitations}" if entry.limitations else "")
        )
    for verdict in manifest.judge:
        lines.append(
            f"Judge, section {verdict.get('section', '?') + 1 if isinstance(verdict.get('section'), int) else '?'}: {verdict.get('verdict')} {verdict.get('notes', '')}".strip()
        )
    return lines


def _template_bytes(docx_key: str | None, open_blob: Any) -> io.BytesIO:
    if docx_key and docx_key.startswith("packaged:"):
        data = (
            resources.files("policy_advisor.artefacts.templates")
            .joinpath(docx_key.split(":", 1)[1])
            .read_bytes()
        )
        return io.BytesIO(data)
    if docx_key and open_blob is not None:
        with open_blob(docx_key) as handle:
            return io.BytesIO(handle.read())
    data = resources.files("policy_advisor.artefacts.templates").joinpath("memo.docx").read_bytes()
    return io.BytesIO(data)


def render_docx(
    outline: Outline,
    manifest: Manifest,
    *,
    docx_key: str | None = None,
    style: dict[str, Any] | None = None,
    open_blob: Any = None,
) -> bytes:
    style = style or {}
    template = DocxTemplate(_template_bytes(docx_key, open_blob))
    counts = manifest.counts()
    context = {
        "title": outline.title,
        "audience": outline.audience,
        "summary": outline.summary,
        "sections": _paragraph_views(outline, manifest),
        "limitations": outline.limitations or ["None recorded by the drafting run."],
        "provenance_summary": "Sources: "
        + ", ".join(f"{k} {v}" for k, v in sorted(counts.items()))
        + f". Generated {manifest.generated_at}.",
        "provenance": _provenance_lines(manifest),
        "footer": style.get("footer", ""),
    }
    template.render(context)
    out = io.BytesIO()
    template.save(out)
    return out.getvalue()


def render_pptx(
    outline: Outline, manifest: Manifest, *, style: dict[str, Any] | None = None
) -> bytes:
    style = style or {}
    deck = Presentation()
    title_layout, content_layout = deck.slide_layouts[0], deck.slide_layouts[1]
    slide = deck.slides.add_slide(title_layout)
    slide.shapes.title.text = outline.title
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = outline.audience or outline.summary[:200]
    for section in _paragraph_views(outline, manifest):
        slide = deck.slides.add_slide(content_layout)
        slide.shapes.title.text = section["heading"]
        body = slide.placeholders[1].text_frame
        body.clear()
        first = True
        for paragraph in section["paragraphs"]:
            run_text = paragraph["text"][:400]
            para = body.paragraphs[0] if first else body.add_paragraph()
            first = False
            para.text = run_text
            para.font.size = Pt(14)
        notes = slide.notes_slide.notes_text_frame
        notes.text = "\n".join(f"- {p['cite'] or p['status']}" for p in section["paragraphs"])
    closing = deck.slides.add_slide(content_layout)
    closing.shapes.title.text = "Limitations and provenance"
    frame = closing.placeholders[1].text_frame
    frame.clear()
    lines = (outline.limitations or ["None recorded."]) + [
        "Sources: " + ", ".join(f"{k} {v}" for k, v in sorted(manifest.counts().items()))
    ]
    for index, line in enumerate(lines):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.text = line[:300]
        para.font.size = Pt(12)
    if style.get("footer"):
        box = closing.shapes.add_textbox(Inches(0.5), Inches(7), Inches(9), Inches(0.4))
        box.text_frame.text = str(style["footer"])
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()
