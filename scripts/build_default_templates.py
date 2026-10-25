"""Write the packaged default memo template (docs/35):
src/policy_advisor/artefacts/templates/memo.docx, a DOCX whose paragraphs
carry docxtpl's Jinja tags. A firm replaces it by uploading its own DOCX
with the same tags to the templates table.

Usage:
  uv run python -m scripts.build_default_templates
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt

OUT = (
    Path(__file__).resolve().parents[1]
    / "src"
    / "policy_advisor"
    / "artefacts"
    / "templates"
    / "memo.docx"
)


def main() -> int:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    document.add_heading("{{ title }}", level=1)
    document.add_paragraph("{{ audience }}")
    document.add_paragraph("{{ summary }}")
    document.add_paragraph("{%p for section in sections %}")
    document.add_heading("{{ section.heading }}", level=2)
    document.add_paragraph("{%p for paragraph in section.paragraphs %}")
    document.add_paragraph(
        "{{ paragraph.text }}{% if paragraph.cite %} [{{ paragraph.cite }}]{% endif %}"
    )
    document.add_paragraph("{%p endfor %}")
    document.add_paragraph("{%p endfor %}")
    document.add_heading("Limitations", level=2)
    document.add_paragraph("{%p for item in limitations %}")
    document.add_paragraph("{{ item }}", style="List Bullet")
    document.add_paragraph("{%p endfor %}")
    document.add_heading("Provenance", level=2)
    document.add_paragraph("{{ provenance_summary }}")
    document.add_paragraph("{%p for entry in provenance %}")
    document.add_paragraph("{{ entry }}", style="List Bullet")
    document.add_paragraph("{%p endfor %}")
    document.add_paragraph("{{ footer }}")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUT)
    print(f"wrote {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
