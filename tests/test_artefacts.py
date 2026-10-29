"""Artefacts (docs/30 slice P3-E, AC-5): the outline parser, the provenance
manifest that rejects evidence cited as authority, both renderers, and a
drafting run on the research loop with a scripted model and a fake judge."""

import io
import json

import pytest
from docx import Document
from pptx import Presentation

from policy_advisor.artefacts import outline as outline_module
from policy_advisor.artefacts.draft import DRAFT_PLAYBOOK, DraftingAgent
from policy_advisor.artefacts.outline import Outline, Paragraph, Section, parse_outline
from policy_advisor.artefacts.provenance import build_manifest
from policy_advisor.artefacts.render import (
    REJECTED_MARKER,
    UNVERIFIED_MARKER,
    render_docx,
    render_pptx,
)
from policy_advisor.artefacts.store import get_artefact, list_templates
from policy_advisor.generation.case_reasoning_models import StepCheckOutcome
from policy_advisor.storage import LocalBlobStore
from tests.test_research_agent import _agent, _chunk

AUTHORITY = _chunk(
    "Order 5 Rule 3",
    "A defence shall be filed within 30 days.",
    trust="authority",
    document="rules.pdf",
    matter="phase1-demo",
)
EVIDENCE = _chunk("Clause 4", "The Buyer shall pay within thirty days.")


def _outline():
    return Outline(
        title="Late defence memo",
        audience="Partner",
        summary="Whether a day-40 defence is out of time.",
        sections=[
            Section(
                heading="Facts",
                paragraphs=[Paragraph(text="The defence was filed on day 40.", kind="instruction")],
            ),
            Section(
                heading="Authority",
                paragraphs=[
                    Paragraph(text="A defence has 30 days.", sources=["Order 5 Rule 3"]),
                    Paragraph(text="Payment is due in thirty days.", sources=["Clause 4"]),
                    Paragraph(text="Lagos says otherwise.", sources=["Order 99 Rule 1"]),
                    Paragraph(text="USCIS agrees.", sources=["https://www.uscis.gov/x"]),
                ],
            ),
        ],
        limitations=["Lagos rules not checked"],
    )


class Recorded:
    url = "https://www.uscis.gov/x"
    title = "USCIS page"
    source_id = "ng-nis"
    trust_kind = "authority"


def test_parse_outline_tolerates_prose_and_fences_and_refuses_junk():
    raw = _outline().model_dump()
    assert (
        parse_outline("Here is the memo.\n```json\n" + json.dumps(raw) + "\n```").title
        == "Late defence memo"
    )
    assert parse_outline("Done. " + json.dumps(raw)).sections[1].paragraphs[0].sources == [
        "Order 5 Rule 3"
    ]
    with pytest.raises(outline_module.OutlineError):
        parse_outline('no json here {"title": }')


def test_manifest_verifies_authority_rejects_evidence_and_marks_the_rest():
    manifest = build_manifest(
        _outline(), [AUTHORITY, EVIDENCE], {"https://www.uscis.gov/x": Recorded()}
    )
    by_pin = {e.pinpoint: e for e in manifest.entries if e.pinpoint}
    assert (
        by_pin["Order 5 Rule 3"].status == "verified"
        and by_pin["Order 5 Rule 3"].instrument == "rules.pdf"
    )
    assert (
        by_pin["Clause 4"].status == "rejected_evidence"
        and by_pin["Clause 4"].trust_kind == "evidence"
    )
    assert by_pin["Order 99 Rule 1"].status == "unverified"
    assert (
        by_pin["https://www.uscis.gov/x"].status == "verified"
        and by_pin["https://www.uscis.gov/x"].source_id == "ng-nis"
    )
    assert any(e.status == "instruction" for e in manifest.entries)
    assert manifest.rejected_paragraphs == {(1, 1)} and manifest.unverified_paragraphs == {(1, 2)}
    assert manifest.counts() == {
        "instruction": 1,
        "verified": 2,
        "rejected_evidence": 1,
        "unverified": 1,
    }


def test_renderers_remove_rejected_paragraphs_and_mark_unverified_ones():
    outline = _outline()
    manifest = build_manifest(
        outline, [AUTHORITY, EVIDENCE], {"https://www.uscis.gov/x": Recorded()}
    )
    manifest.judge.append({"section": 1, "verdict": "FLAGGED", "notes": "one inference too far"})
    docx = Document(
        io.BytesIO(render_docx(outline, manifest, style={"footer": "Review before use."}))
    )
    text = "\n".join(p.text for p in docx.paragraphs)
    assert "Late defence memo" in text and "A defence has 30 days. [Order 5 Rule 3]" in text
    assert REJECTED_MARKER in text and "Payment is due in thirty days." not in text
    assert "Lagos says otherwise." + UNVERIFIED_MARKER in text
    assert "Judge, section 2: FLAGGED" in text and "Review before use." in text
    deck = Presentation(io.BytesIO(render_pptx(outline, manifest)))
    titles = [s.shapes.title.text for s in deck.slides]
    assert titles == ["Late defence memo", "Facts", "Authority", "Limitations and provenance"]
    authority_slide = deck.slides[2]
    body = "\n".join(p.text for p in authority_slide.placeholders[1].text_frame.paragraphs)
    assert REJECTED_MARKER in body and "Payment is due" not in body
    assert "Order 5 Rule 3" in authority_slide.notes_slide.notes_text_frame.text


def test_drafting_run_is_authority_only_and_stores_files_and_manifest(fresh_db, tmp_path):
    async def script(tools, system, prompt):
        assert "search_authority" in tools and "search_firm_knowledge" in tools
        assert "search_case_documents" not in tools and "analyze_case" not in tools
        assert (
            "Playbook: Drafting (authority only)" in system
            and "call the submit_outline tool" in prompt
        )
        found = await tools["search_authority"].call({"question": "time to file a defence"})
        assert "Order 5 Rule 3" in found
        none = await tools["search_firm_knowledge"].call({"question": "precedent memo"})
        assert "No collections" in none
        return "Drafted. " + json.dumps(_outline().model_dump()), {
            "input_tokens": 50,
            "output_tokens": 20,
        }

    judged = []

    def fake_judge(text, chunks):
        judged.append((text, [c.metadata["locator"] for c in chunks]))
        return StepCheckOutcome(
            faithful=False, judge_flagged=True, judge_notes="stretches the rule"
        )

    blobs = LocalBlobStore(tmp_path / "blobs")
    agent = DraftingAgent(_agent(script), blob_store=blobs, judge=fake_judge)
    from policy_advisor.ingestion import matter_store

    matter_store.set_matter_owner("case-1", "alice")
    result = agent.draft(
        "Memo for the partner on a day-40 defence.",
        matter_id="case-1",
        username="alice",
        kind="both",
    )
    assert result.artefact.status == "draft" and result.artefact.kind == "both"
    assert (
        result.manifest.counts()["verified"] >= 1
        and result.manifest.counts()["rejected_evidence"] == 1
    )
    assert judged and judged[0][1] == [
        "Order 5 Rule 3"
    ]  # only the Authority section is judged, with its passages
    assert result.manifest.judge == [
        {"section": 1, "verdict": "FLAGGED", "notes": "stretches the rule"}
    ]
    stored = get_artefact(result.artefact.artefact_id)
    assert (
        stored.docx_key
        and stored.pptx_key
        and blobs.exists(stored.docx_key)
        and blobs.exists(stored.pptx_key)
    )
    assert stored.manifest["counts"]["rejected_evidence"] == 1 and stored.run_id
    assert stored.template_id == "memo-default"
    assert DRAFT_PLAYBOOK.allows_tool("search_authority") and not DRAFT_PLAYBOOK.allows_tool(
        "search_case_documents"
    )
    assert [t.template_id for t in list_templates()] == ["memo-default", "slides-default"]


def _drafter(script, tmp_path, *, max_calls=12):
    from policy_advisor.ingestion import matter_store

    matter_store.set_matter_owner("case-1", "alice")

    def judge(text, chunks):
        return StepCheckOutcome(faithful=True, judge_flagged=False, judge_notes="")

    return DraftingAgent(
        _agent(script, max_calls=max_calls),
        blob_store=LocalBlobStore(tmp_path / "blobs"),
        judge=judge,
    )


def test_the_outline_handed_in_through_the_tool_is_used_and_a_rejected_one_is_retried(
    fresh_db, tmp_path
):
    # The first nightly run (2026-10-28) ended all three drafts without a
    # valid outline when the outline was JSON written into the reply.
    replies = []

    async def script(tools, system, prompt):
        await tools["search_authority"].call({"question": "time to file a defence"})
        bad = _outline().model_dump()
        bad["sections"][0]["paragraphs"][0]["kind"] = "evidence"
        replies.append(await tools["submit_outline"].call(bad))
        good = _outline().model_dump()
        good["sections"] = json.dumps(good["sections"])  # a list sent as a JSON string
        replies.append(await tools["submit_outline"].call(good))
        return "Outline submitted.", {"input_tokens": 50, "output_tokens": 5}

    result = _drafter(script, tmp_path).draft(
        "Memo on a day-40 defence.", matter_id="case-1", username="alice", kind="memo"
    )
    assert replies[0].startswith("Outline rejected") and "kind" in replies[0]
    assert replies[1].startswith("Outline recorded")
    assert result.outline == _outline()
    assert get_artefact(result.artefact.artefact_id).docx_key


def test_the_outline_can_be_handed_in_after_the_tool_budget_is_spent(fresh_db, tmp_path):
    async def script(tools, system, prompt):
        await tools["search_authority"].call({"question": "time to file a defence"})
        refused = await tools["search_authority"].call({"question": "anything else"})
        assert "budget exhausted" in refused
        recorded = await tools["submit_outline"].call(_outline().model_dump())
        assert recorded.startswith("Outline recorded")
        return "Done.", {}

    result = _drafter(script, tmp_path, max_calls=1).draft(
        "Memo on a day-40 defence.", matter_id="case-1", username="alice", kind="memo"
    )
    assert result.outline.title == _outline().title


def test_a_run_with_no_outline_says_why_without_quoting_the_reply(fresh_db, tmp_path):
    async def script(tools, system, prompt):
        return (
            'Here it is: {"title": "Memo", "sections": [{"heading": "A", "paragraphs": [{"text": "He said "no" twice"}]}]}',
            {},
        )

    with pytest.raises(outline_module.OutlineError) as caught:
        _drafter(script, tmp_path).draft("Memo.", matter_id="case-1", username="alice", kind="memo")
    message = str(caught.value)
    assert "JSON error at character" in message and "He said" not in message


@pytest.mark.parametrize("kind", ["memo", "both", "slides"])
def test_a_memo_run_names_the_playbook_sections_in_the_outline_instruction(
    fresh_db, tmp_path, kind
):
    # The third nightly run (docs/39 §9a): the headings reached only the
    # system prompt, and both memos used topic headings instead.
    prompts = []

    async def script(tools, system, prompt):
        prompts.append(prompt)
        await tools["submit_outline"].call(_outline().model_dump())
        return "Done.", {}

    _drafter(script, tmp_path).draft("Memo.", matter_id="case-1", username="alice", kind=kind)
    named = "these headings: " + "; ".join(DRAFT_PLAYBOOK.output_sections)
    if kind == "slides":
        assert named not in prompts[0], "a deck's sections await the lawyer review"
    else:
        assert named in prompts[0] and prompts[0].index("submit_outline") < prompts[0].index(named)
