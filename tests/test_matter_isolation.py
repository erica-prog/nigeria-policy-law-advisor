"""Integration test for the client-confidentiality boundary added in Phase A:
a query scoped to one matter must never return another matter's chunks, even
when both matters contain a document with the exact same filename and even
near-identical text. Uses the real embedding model and persisted Chroma
store (no mocks) because this guarantee is the one thing in this phase that
must not be wrong - worth the extra runtime."""

import uuid

import pytest
from docx import Document as DocxDocument

from policy_advisor.ingestion.ingest_document import add_document, remove_document
from policy_advisor.retrieval.hybrid_retriever import HybridRetriever


@pytest.fixture
def two_isolated_matters(tmp_path):
    matter_a = f"test-matter-a-{uuid.uuid4().hex[:8]}"
    matter_b = f"test-matter-b-{uuid.uuid4().hex[:8]}"

    # Same filename in both matters - exercises that chunk_id prefixing by
    # matter_id (not filename alone) is what actually prevents collisions,
    # not an accidental side effect of the two filenames differing.
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    doc_a = tmp_path / "a" / "contract.docx"
    doc_b = tmp_path / "b" / "contract.docx"

    docx_doc = DocxDocument()
    docx_doc.add_paragraph("This contract between Acme Corp and Foxglove Ltd is confidential to matter A.")
    docx_doc.add_paragraph("The exclusivity clause runs for eighteen months from signature.")
    docx_doc.save(doc_a)

    docx_doc_b = DocxDocument()
    docx_doc_b.add_paragraph("This contract between Acme Corp and Foxglove Ltd is confidential to matter A.")
    docx_doc_b.add_paragraph("The exclusivity clause runs for eighteen months from signature.")
    docx_doc_b.save(doc_b)

    add_document(matter_a, doc_a)
    add_document(matter_b, doc_b)

    yield matter_a, matter_b

    remove_document(matter_a, "contract.docx")
    remove_document(matter_b, "contract.docx")


def test_query_in_one_matter_never_returns_another_matters_chunks(two_isolated_matters):
    matter_a, matter_b = two_isolated_matters
    retriever = HybridRetriever()

    results_a = retriever.retrieve("How long does the exclusivity clause run for?", top_k=5, matter_id=matter_a)
    results_b = retriever.retrieve("How long does the exclusivity clause run for?", top_k=5, matter_id=matter_b)

    assert results_a, "expected matter A's own document to be retrieved"
    assert results_b, "expected matter B's own document to be retrieved"

    for chunk in results_a:
        assert chunk.metadata["matter_id"] == matter_a
    for chunk in results_b:
        assert chunk.metadata["matter_id"] == matter_b
